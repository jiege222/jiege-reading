"""Generate the Word plan from Markdown; requires the python-docx package."""

from datetime import datetime, timezone
from hashlib import sha256
from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docx" / "技术方案-MVP.md"
OUTPUT = SOURCE.with_suffix(".docx")


def set_font(style, size, color=None):
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(size)
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if color:
        style.font.color.rgb = RGBColor.from_string(color)


def build():
    # Git may check Markdown out as CRLF on Windows; line endings are not content.
    source_bytes = SOURCE.read_bytes().replace(b"\r\n", b"\n")
    lines = source_bytes.decode("utf-8").splitlines()
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)

    set_font(document.styles["Normal"], 10.5, "253046")
    normal = document.styles["Normal"].paragraph_format
    normal.line_spacing = 1.25
    normal.space_after = Pt(7)
    for style, size in [("Title", 23), ("Heading 1", 15), ("Heading 2", 12)]:
        set_font(document.styles[style], size, "163B60")

    header = section.header.paragraphs[0]
    header.text = "杰哥阅读 · MVP 技术方案"
    set_font(document.styles["Header"], 9, "66758A")
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")

    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", cell) for cell in cells):
                    rows.append(cells)
                index += 1
            table = document.add_table(rows=0, cols=len(rows[0]))
            table.style = "Table Grid"
            for row_index, values in enumerate(rows):
                row = table.add_row()
                for cell, value in zip(row.cells, values):
                    cell.text = value
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(4)
                        paragraph.paragraph_format.space_before = Pt(4)
                        for run in paragraph.runs:
                            run.font.size = Pt(9)
                            if row_index == 0:
                                run.bold = True
                    if row_index == 0:
                        shading = OxmlElement("w:shd")
                        shading.set(qn("w:fill"), "E7EFF7")
                        cell._tc.get_or_add_tcPr().append(shading)
                if row_index == 0:
                    repeat = OxmlElement("w:tblHeader")
                    row._tr.get_or_add_trPr().append(repeat)
            document.add_paragraph()
            continue
        if line.startswith("# "):
            document.add_paragraph(line[2:], style="Title")
        elif line.startswith("## "):
            document.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            document.add_heading(line[4:], level=2)
        else:
            document.add_paragraph(line)
        index += 1

    document.core_properties.title = "杰哥阅读：MVP 技术方案"
    document.core_properties.subject = "React + TypeScript + Vite + IndexedDB"
    document.core_properties.author = "杰哥阅读"
    document.core_properties.created = datetime.now(timezone.utc)
    document.core_properties.modified = document.core_properties.created
    document.core_properties.keywords = "source-sha256:" + sha256(source_bytes).hexdigest()
    document.save(OUTPUT)
    print(f"Generated: {OUTPUT}")


if __name__ == "__main__":
    build()
